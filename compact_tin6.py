from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

SOURCE = Path(r"D:\KNTT-Tin 6.docx")
OUTPUT = Path(r"D:\ERG\erg-spring\KNTT-Tin 6_da_rut_gon.docx")

# Include the date block, topic title, and its first lesson only.
KEEP_PARAGRAPH_RANGES = [
    (4, 254),       # Chủ đề A – Bài 1
    (1322, 1504),   # Chủ đề B – Bài 1
    (2104, 2292),   # Chủ đề C – Bài 1
    (3306, 3480),   # Chủ đề D – Bài 1
    (3915, 4148),   # Chủ đề E – Bài 1
    (5576, 5829),   # Chủ đề F – Bài 1
]

W_BR = qn("w:br")
W_SECTPR = qn("w:sectPr")


def in_kept_range(index):
    return any(start <= index < end for start, end in KEEP_PARAGRAPH_RANGES)


def is_effectively_empty(paragraph):
    if paragraph.text.strip():
        return False
    element = paragraph._element
    # Preserve section/page-break carrier paragraphs; other empty paragraphs
    # are surplus Enter lines.
    return not (element.find(W_SECTPR) is not None or element.findall(".//" + W_BR))


doc = Document(SOURCE)
body = doc._element.body
paragraph_index = 0

# Remove blocks outside the chosen six first-lesson ranges. Tables are kept
# according to the paragraph position immediately before them in document order.
for child in list(body):
    tag = child.tag
    if tag == qn("w:p"):
        keep = in_kept_range(paragraph_index)
        paragraph_index += 1
    elif tag == qn("w:tbl"):
        keep = in_kept_range(paragraph_index - 1)
    elif tag == W_SECTPR:
        keep = True
    else:
        keep = False
    if not keep:
        body.remove(child)

# Remove surplus blank paragraphs after the content selection.
for paragraph in list(doc.paragraphs):
    if is_effectively_empty(paragraph):
        paragraph._element.getparent().remove(paragraph._element)

# Use the smallest practical Word page margins and compact header/footer bounds.
for section in doc.sections:
    section.top_margin = Inches(0.1)
    section.bottom_margin = Inches(0.1)
    section.left_margin = Inches(0.1)
    section.right_margin = Inches(0.1)
    section.header_distance = Inches(0.05)
    section.footer_distance = Inches(0.05)

def format_paragraph(paragraph):
    pf = paragraph.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.line_spacing = 1.0
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    for run in paragraph.runs:
        run.font.size = Pt(12)

def format_container(paragraphs, tables):
    for paragraph in paragraphs:
        format_paragraph(paragraph)
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                format_container(cell.paragraphs, cell.tables)

# Apply 12 pt / single-spacing to body, headers, footers, and all table cells.
format_container(doc.paragraphs, doc.tables)
for section in doc.sections:
    for part in (section.header, section.footer):
        format_container(part.paragraphs, part.tables)

# Make style defaults match direct formatting so any future inserted text follows it.
for style in doc.styles:
    if style.type == 1:  # paragraph styles
        style.font.size = Pt(12)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)

doc.save(OUTPUT)
print(OUTPUT)
