from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Twips

SOURCE = Path(r"D:\ERG\erg-spring\KNTT-Tin 6_da_rut_gon.docx")
OUTPUT = Path(r"D:\ERG\erg-spring\KNTT-Tin 6_tiet_kiem_giay.docx")

W_BR = qn("w:br")
W_COLS = qn("w:cols")
W_TBLW = qn("w:tblW")
W_TBLIND = qn("w:tblInd")
W_TBLLAYOUT = qn("w:tblLayout")
W_GRIDCOL = qn("w:gridCol")
W_TCW = qn("w:tcW")


def set_dxa(element, value, kind="dxa"):
    element.set(qn("w:w"), str(value))
    element.set(qn("w:type"), kind)


def child_or_add(parent, tag):
    child = parent.find(tag)
    if child is None:
        child = OxmlElement("w:" + tag.rsplit("}", 1)[-1])
        parent.append(child)
    return child


def remove_surplus_empty_paragraphs(paragraphs):
    for paragraph in list(paragraphs):
        if paragraph.text.strip():
            continue
        has_break = paragraph._element.findall(".//" + W_BR)
        has_section = paragraph._element.find(qn("w:sectPr")) is not None
        if not has_break and not has_section:
            parent = paragraph._element.getparent()
            parent.remove(paragraph._element)


def set_table_width(table, usable_dxa):
    tbl_pr = table._tbl.tblPr
    set_dxa(child_or_add(tbl_pr, W_TBLW), usable_dxa)
    set_dxa(child_or_add(tbl_pr, W_TBLIND), 0)
    layout = child_or_add(tbl_pr, W_TBLLAYOUT)
    layout.set(qn("w:type"), "fixed")
    table.autofit = False

    grid = table._tbl.tblGrid
    grid_cols = list(grid.findall(W_GRIDCOL))
    old_widths = [int(col.get(qn("w:w"), "0")) for col in grid_cols]
    if not old_widths or not any(old_widths):
        return
    scale = usable_dxa / sum(old_widths)
    widths = [round(width * scale) for width in old_widths]
    widths[-1] += usable_dxa - sum(widths)
    for col, width in zip(grid_cols, widths):
        col.set(qn("w:w"), str(width))

    # Apply the scaled widths to cells. This preserves each table's original
    # column proportions while allowing it to reach the right page margin.
    for row in table.rows:
        col_index = 0
        for cell in row.cells:
            span = cell._tc.tcPr.find(qn("w:gridSpan"))
            span_count = int(span.get(qn("w:val"))) if span is not None else 1
            width = sum(widths[col_index:col_index + span_count])
            tcw = child_or_add(cell._tc.tcPr, W_TCW)
            set_dxa(tcw, width)
            cell.width = Twips(width)
            col_index += span_count
        for row_cell in row.cells:
            for nested in row_cell.tables:
                set_table_width(nested, usable_dxa)


doc = Document(SOURCE)

# Use the complete printable width, rather than the inherited narrow 6.67-inch
# text column that caused unused space on the right side of each page.
for section in doc.sections:
    sect_pr = section._sectPr
    cols = sect_pr.find(W_COLS)
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    for child in list(cols):
        cols.remove(child)
    cols.set(qn("w:num"), "1")
    cols.set(qn("w:space"), "0")
    usable_dxa = section.page_width.twips - section.left_margin.twips - section.right_margin.twips

    for table in doc.tables:
        set_table_width(table, usable_dxa)

# Column breaks from the original multi-column document create partly empty
# pages after the document is converted to a single wide column.
for paragraph in doc.paragraphs:
    for br in list(paragraph._element.findall(".//" + W_BR)):
        if br.get(qn("w:type")) == "column":
            br.getparent().remove(br)

remove_surplus_empty_paragraphs(doc.paragraphs)
doc.save(OUTPUT)
print(OUTPUT)
