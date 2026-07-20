import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.shared import Inches, Pt

SOURCE = Path(r"D:\ERG\erg-spring\tin9_selected_pages.json")
OUTPUT = Path(r"D:\ERG\erg-spring\Tin 9 - mot bai moi chu de.docx")

TOPIC_STARTS = {
    1: "CHỦ ĐỀ 1 — BÀI 1: THẾ GIỚI KĨ THUẬT SỐ",
    15: "CHỦ ĐỀ 2 — BÀI 2: THÔNG TIN TRONG GIẢI QUYẾT VẤN ĐỀ",
    37: "CHỦ ĐỀ 3 — BÀI 4: MỘT SỐ VẤN ĐỀ PHÁP LÍ VỀ SỬ DỤNG DỊCH VỤ INTERNET",
    53: "CHỦ ĐỀ 4 — BÀI 5: TÌM HIỂU PHẦN MỀM MÔ PHỎNG",
    216: "CHỦ ĐỀ 5 — BÀI 14: GIẢI QUYẾT VẤN ĐỀ",
    253: "CHỦ ĐỀ 6 — BÀI 17: TIN HỌC VÀ THẾ GIỚI NGHỀ NGHIỆP",
}

pages = json.loads(SOURCE.read_text(encoding="utf-8"))
doc = Document()
for section in doc.sections:
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
normal.paragraph_format.space_after = Pt(0)

for page_index, record in enumerate(pages):
    source_page = record["page"]
    if source_page in TOPIC_STARTS:
        heading = doc.add_paragraph()
        heading.alignment = 1
        run = heading.add_run(TOPIC_STARTS[source_page])
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(13)
        heading.paragraph_format.space_after = Pt(4)

    # The text is taken from the visible PDF page. Keep every non-empty line,
    # while avoiding visually redundant blank runs from the PDF viewer.
    lines = [line.strip() for line in record["text"].splitlines() if line.strip()]
    for line in lines:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(line)
        run.font.name = "Arial"
        run.font.size = Pt(11)

    if page_index != len(pages) - 1:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

doc.save(OUTPUT)
print(OUTPUT)
